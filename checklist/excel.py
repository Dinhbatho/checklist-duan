"""Xuất kết quả ra file .xlsx — giữ nguyên bố cục của CheckList_Dự_án_đầu_tư_v4.xlsm.

Ba sheet: Info | ListFolder | ListFile, đúng cột, đúng màu, đúng quy ước như bản VBA.
"""

from __future__ import annotations

import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .scanner import KetQuaQuet

# ---- Màu mặc định (giống Info!F5:F9 của bản VBA) --------------------------- #
MAU_SUA = "FFFF00"      # File đã thay đổi nội dung
MAU_MOI = "C8C8C8"      # File được tạo mới
MAU_F1 = "9BC2E6"       # Folder CÓ chứa file hợp lệ
MAU_F0 = "C6E0BE"       # Folder KHÔNG chứa file hợp lệ

TXT_MOI = "File được tạo mới"
TXT_SUA = "File đã thay đổi nội dung"

# Màu luân phiên cho từng cấp đường dẫn (giống Sub ToMauDuongDan)
MAU_CAP = ["0000FF", "009900", "FF6600", "990099", "C80000"]

DINH_DANG_NGAY = "d/m/yy h:mm AM/PM"
VIEN = Border(*[Side(style="thin", color="FF000000")] * 4)


def _rut_gon(p: str, goc: str) -> str:
    """Bỏ phần thư mục cha của root khỏi đường dẫn hiển thị (giống Function RutGon)."""
    if goc and p.lower().startswith(goc.lower()):
        return p[len(goc):]
    return p


def _cong_thuc_link(p: str, hien_thi: str) -> str:
    """HYPERLINK chỉ chạy khi địa chỉ <= 255 ký tự và không chứa '#'."""
    if len(p) <= 255 and "#" not in p:
        an_toan = p.replace('"', '""')
        return f'=HYPERLINK("{an_toan}","{hien_thi}")'
    return f"{hien_thi} (mở bằng cột H)"


def _rich_text_duong_dan(chuoi: str):
    """Tô màu từng cấp thư mục. Trả về CellRichText nếu openpyxl hỗ trợ, ngược lại None."""
    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
    except ImportError:
        return None
    khoi = []
    phan = re.split(r"(?<=[\\/])", chuoi)
    for i, van_ban in enumerate(phan):
        if not van_ban:
            continue
        khoi.append(TextBlock(InlineFont(color=MAU_CAP[i % len(MAU_CAP)], b=True), van_ban))
    return CellRichText(khoi) if khoi else None


class BaoCaoExcel:
    def __init__(self, kq: KetQuaQuet, root: str, duoi_file: list[str],
                 lan_quet_truoc: datetime | None, chay_lan_dau: bool,
                 to_mau_duong_dan: bool = True) -> None:
        self.kq = kq
        self.root = root
        self.duoi_file = duoi_file
        self.lan_quet_truoc = lan_quet_truoc
        self.chay_lan_dau = chay_lan_dau
        self.to_mau_duong_dan = to_mau_duong_dan
        r = root.rstrip("\\/")
        i = max(r.rfind("\\"), r.rfind("/"))
        self.goc_cha = r[:i] if i > 0 else ""

    # ------------------------------------------------------------------ #
    def _trang_thai(self, ngay_tao: datetime | None, ngay_sua: datetime | None) -> str:
        """Cột 'Ghi chú': file mới tạo / đã sửa so với lần quét trước."""
        if self.chay_lan_dau or not self.lan_quet_truoc:
            return ""
        if ngay_tao and self.lan_quet_truoc < ngay_tao:
            return TXT_MOI
        if ngay_sua and self.lan_quet_truoc < ngay_sua:
            return TXT_SUA
        return ""

    # ------------------------------------------------------------------ #
    def xuat(self, duong_dan_ra: str) -> str:
        wb = Workbook()
        self._sheet_info(wb.active)
        self._sheet_list_folder(wb.create_sheet("ListFolder"))
        self._sheet_list_file(wb.create_sheet("ListFile"))
        wb.active = wb["ListFile"]
        wb.save(duong_dan_ra)
        return duong_dan_ra

    # ------------------------------------------------------------------ #
    def _sheet_info(self, ws: Worksheet) -> None:
        ws.title = "Info"
        dam = Font(bold=True)

        ws["B1"] = "Folder gốc:"; ws["B1"].font = dam
        ws["C1"] = self.root
        ws["E2"] = "Số lượng Folder:"; ws["F2"] = len(self.kq.folders)
        ws["G2"] = "Số lượng File:";   ws["H2"] = self.kq.so_file
        ws["E3"] = "Thời gian quét:";  ws["F3"] = datetime.now()
        ws["F3"].number_format = "dd/mm/yyyy hh:mm:ss"
        ws["E4"] = "Lần quét trước:"
        if self.lan_quet_truoc:
            ws["F4"] = self.lan_quet_truoc
            ws["F4"].number_format = "dd/mm/yyyy hh:mm:ss"
        ws["E5"] = "Thời gian chạy:";  ws["F5"] = f"{self.kq.giay:.1f} giây"
        ws["E6"] = "Số lỗi:";          ws["F6"] = len(self.kq.loi)
        if self.kq.bi_huy:
            ws["E7"] = "Trạng thái:";  ws["F7"] = "ĐÃ DỪNG GIỮA CHỪNG (Ctrl+C)"
        elif self.kq.vuot_gioi_han:
            ws["E7"] = "Trạng thái:";  ws["F7"] = "VƯỢT GIỚI HẠN DÒNG EXCEL — ĐÃ CẮT BỚT"

        # Bảng đuôi file
        ws["A3"] = "STT"; ws["B3"] = "Đuôi File cần thống kê"
        for c in ("A3", "B3"):
            ws[c].font = dam
        for i, d in enumerate(self.duoi_file, start=1):
            ws.cell(row=3 + i, column=1, value=i)
            ws.cell(row=3 + i, column=2, value=d)

        # Chú giải màu (giữ đúng vị trí E5:F9 của bản cũ -> dời xuống E10 cho gọn)
        ws["E10"] = "NỘI DUNG"; ws["F10"] = "MẦU"
        ws["E10"].font = dam; ws["F10"].font = dam
        chu_giai = [(TXT_SUA, MAU_SUA), (TXT_MOI, MAU_MOI), ("", None),
                    ("Folder chứa File hợp lệ", MAU_F1),
                    ("Folder KHÔNG chứa File hợp lệ", MAU_F0)]
        for i, (txt, mau) in enumerate(chu_giai, start=11):
            ws.cell(row=i, column=5, value=txt)
            if mau:
                ws.cell(row=i, column=6).fill = PatternFill("solid", fgColor=mau)

        # Log lỗi — cột N:R như bản VBA
        tieu_de = ["Mã lỗi", "Mô tả", "Hàm", "Thời điểm", "Đường dẫn"]
        for j, t in enumerate(tieu_de, start=14):
            o = ws.cell(row=1, column=j, value=t)
            o.font = dam
        for i, l in enumerate(self.kq.loi, start=2):
            ws.cell(row=i, column=14, value=l.ma_loi)
            ws.cell(row=i, column=15, value=l.mo_ta)
            ws.cell(row=i, column=16, value=l.ham)
            o = ws.cell(row=i, column=17, value=l.thoi_diem)
            o.number_format = "dd/mm/yyyy hh:mm:ss"
            ws.cell(row=i, column=18, value=l.duong_dan)

        for col, w in {"A": 6, "B": 24, "C": 60, "E": 30, "F": 22, "G": 16, "H": 12,
                       "N": 9, "O": 46, "P": 18, "Q": 20, "R": 70}.items():
            ws.column_dimensions[col].width = w

    # ------------------------------------------------------------------ #
    def _sheet_list_folder(self, ws: Worksheet) -> None:
        ws["A1"] = len(self.kq.folders)
        for i, f in enumerate(self.kq.folders):
            ws.cell(row=i + 2, column=1, value=i)
            o = ws.cell(row=i + 2, column=2, value=f.duong_dan)
            o.number_format = "@"
            ws.cell(row=i + 2, column=3, value=f.ma)
        ws.column_dimensions["B"].width = 120
        ws.column_dimensions["C"].width = 25
        ws.freeze_panes = "A2"

    # ------------------------------------------------------------------ #
    def _sheet_list_file(self, ws: Worksheet) -> None:
        tieu_de = ["STT", "Nội dung", "Ghi chú", "Mở File", "Ngày tạo", "Ngày sửa Chữa", ""]
        for j, t in enumerate(tieu_de, start=1):
            o = ws.cell(row=1, column=j, value=t)
            o.font = Font(bold=True)
        ws.cell(row=1, column=8, value="Đường dẫn đầy đủ")
        ws.cell(row=1, column=9, value="Cờ")

        dong = 2
        stt_file = 0
        dong_folder: list[int] = []

        for f in self.kq.folders:
            # --- dòng tiêu đề folder ---
            ws.cell(row=dong, column=1, value=_cong_thuc_link(f.duong_dan, f.ma))
            hien_thi = _rut_gon(f.duong_dan, self.goc_cha)
            o = ws.cell(row=dong, column=2)
            rt = _rich_text_duong_dan(hien_thi) if self.to_mau_duong_dan else None
            o.value = rt if rt is not None else hien_thi
            if rt is None:
                o.font = Font(bold=True)
            o.number_format = "@"
            ws.cell(row=dong, column=8, value=f.duong_dan).number_format = "@"
            ws.cell(row=dong, column=9, value="F1" if f.files else "F0")
            dong_folder.append(dong)
            dong += 1

            # --- các file trong folder ---
            for fl in f.files:
                stt_file += 1
                ws.cell(row=dong, column=1, value=stt_file)
                ws.cell(row=dong, column=2, value=fl.ten).number_format = "@"
                tt = self._trang_thai(fl.ngay_tao, fl.ngay_sua)
                if tt:
                    ws.cell(row=dong, column=3, value=tt)
                ws.cell(row=dong, column=4, value=_cong_thuc_link(fl.duong_dan, "Open File"))
                if fl.ngay_tao:
                    o = ws.cell(row=dong, column=5, value=fl.ngay_tao)
                    o.number_format = DINH_DANG_NGAY
                if fl.ngay_sua:
                    o = ws.cell(row=dong, column=6, value=fl.ngay_sua)
                    o.number_format = DINH_DANG_NGAY
                ws.cell(row=dong, column=8, value=fl.duong_dan).number_format = "@"
                dong += 1

        cuoi = dong - 1
        if cuoi < 2:
            return

        # viền + xuống dòng
        for r in range(2, cuoi + 1):
            for c in range(1, 8):
                ws.cell(row=r, column=c).border = VIEN
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

        # Conditional formatting — đúng thứ tự ưu tiên của bản VBA
        vung = f"A2:G{cuoi}"
        vung_b = f"B2:G{cuoi}"
        ws.conditional_formatting.add(vung, FormulaRule(
            formula=['$I2="F1"'], fill=PatternFill("solid", bgColor=MAU_F1),
            font=Font(bold=True), stopIfTrue=True))
        ws.conditional_formatting.add(vung, FormulaRule(
            formula=['$I2="F0"'], fill=PatternFill("solid", bgColor=MAU_F0),
            font=Font(bold=True), stopIfTrue=True))
        ws.conditional_formatting.add(vung_b, FormulaRule(
            formula=[f'$C2="{TXT_MOI}"'], fill=PatternFill("solid", bgColor=MAU_MOI),
            stopIfTrue=True))
        ws.conditional_formatting.add(vung_b, FormulaRule(
            formula=[f'$C2="{TXT_SUA}"'], fill=PatternFill("solid", bgColor=MAU_SUA),
            stopIfTrue=True))

        ws.auto_filter.ref = f"A1:G{cuoi}"
        ws.freeze_panes = "A2"

        for col, w in {"A": 27, "B": 63.6, "C": 25.5, "D": 15, "E": 19.1,
                       "F": 18.25, "G": 56.5, "H": 19.25, "I": 68}.items():
            ws.column_dimensions[col].width = w
        for col in ("H", "I"):
            ws.column_dimensions[col].hidden = True
        _ = get_column_letter  # giữ import cho tiện mở rộng
